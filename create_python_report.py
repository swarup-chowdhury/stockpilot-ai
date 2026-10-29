from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'reports'
ASSETS = OUT / 'assets'
OUT.mkdir(exist_ok=True)

NAVY = '10233F'
INDIGO = '5D4BF1'
TEAL = '0B9E9A'
LIGHT = 'F3F6FB'
MID = '667085'
DARK = '172033'
WHITE = 'FFFFFF'
BORDER = 'D9E1EE'


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def cell_border(cell, color=BORDER, size='6'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def set_cell_text(cell, text, bold=False, color=DARK, size=9.2):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = 'Aptos'
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('StockPilot AI — Python Programming Project Report  |  ')
    run.font.name = 'Aptos'
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(fld)


def add_rule(paragraph, color=INDIGO):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '18')
    bottom.set(qn('w:space'), '5')
    bottom.set(qn('w:color'), color)
    pbdr.append(bottom)
    pPr.append(pbdr)


def add_heading(doc, text, level=1, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.name = 'Aptos Display'
    r.bold = True
    r.font.size = Pt(20 if level == 1 else 13)
    r.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else INDIGO)
    if level == 1:
        add_rule(p)
    if subtitle:
        sp = doc.add_paragraph(subtitle)
        sp.style = 'Caption'
        sp.paragraph_format.space_after = Pt(8)
    return p


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.first_line_indent = Inches(0.22)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(NAVY)
        rest = text[len(bold_prefix):]
        p.add_run(rest)
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.05
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.05
        p.add_run(item)


def add_table(doc, headers, rows, widths=None, font_size=8.6):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, color=WHITE, size=font_size)
        shade(hdr[i], NAVY)
        cell_border(hdr[i], NAVY)
        set_cell_margins(hdr[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
            shade(cells[i], 'FFFFFF' if ridx % 2 == 0 else LIGHT)
            cell_border(cells[i])
            set_cell_margins(cells[i])
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout(doc, title, text, color=TEAL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, 'EFFAF9' if color == TEAL else 'F3F1FF')
    cell_border(cell, color, '10')
    set_cell_margins(cell, top=130, start=160, bottom=130, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(color)
    p2 = cell.add_paragraph(text)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.05
    for r in p2.runs:
        r.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code(doc, code):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, '162033')
    cell_border(cell, '162033')
    set_cell_margins(cell, top=140, start=170, bottom=140, end=170)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for i, line in enumerate(code.strip('\n').splitlines()):
        r = p.add_run(line + ('\n' if i < len(code.strip('\n').splitlines()) - 1 else ''))
        r.font.name = 'Consolas'
        r.font.size = Pt(8.2)
        r.font.color.rgb = RGBColor.from_string('E4ECFF')
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.italic = True
    r.font.name = 'Aptos'
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(MID)


def page_break(doc):
    doc.add_page_break()


def configure(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.3)
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    for name, size, color in [('Title', 30, NAVY), ('Heading 1', 20, NAVY), ('Heading 2', 13, INDIGO), ('Heading 3', 11, TEAL)]:
        style = styles[name]
        style.font.name = 'Aptos Display'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
    if 'Caption' in styles:
        styles['Caption'].font.name = 'Aptos'
        styles['Caption'].font.size = Pt(8.5)
        styles['Caption'].font.italic = True
        styles['Caption'].font.color.rgb = RGBColor.from_string(MID)
    footer = sec.footer.paragraphs[0]
    add_page_number(footer)
    header = sec.header.paragraphs[0]
    header.text = 'STOCKPILOT AI  /  PYTHON PROGRAMMING'
    header.runs[0].font.name = 'Aptos'
    header.runs[0].font.size = Pt(7.5)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(MID)


def build():
    doc = Document()
    configure(doc)

    # Cover page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run('STOCKPILOT AI')
    r.font.name = 'Aptos Display'; r.font.size = Pt(28); r.bold = True; r.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run('Smart Inventory & Demand Intelligence System')
    r.font.size = Pt(13); r.bold = True; r.font.color.rgb = RGBColor.from_string(INDIGO)
    doc.add_paragraph()
    cover_line = doc.add_paragraph()
    cover_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_line.add_run('PYTHON PROGRAMMING PROJECT REPORT').bold = True
    cover_line.runs[0].font.size = Pt(14); cover_line.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    doc.add_paragraph(); doc.add_paragraph()
    for label, value in [
        ('Submitted by', 'SWARUP CHOWDHURY'),
        ('Roll Number', '26MCI10229'),
        ('Section', '26MAM4(B)'),
    ]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(4)
        rr = p.add_run(label.upper() + '\n' + value)
        rr.font.size = Pt(12 if label == 'Submitted by' else 10.5); rr.bold = True if label == 'Submitted by' else False; rr.font.color.rgb = RGBColor.from_string(DARK)
    doc.add_paragraph(); doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('in partial fulfillment for the award of the degree of')
    r.font.size = Pt(10); r.font.color.rgb = RGBColor.from_string(MID)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Master of Computer Applications in Artificial Intelligence and\nMachine Learning (MCA-AIML)')
    r.bold = True; r.font.size = Pt(12); r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Subject: Python Programming')
    r.font.size = Pt(11); r.bold = True; r.font.color.rgb = RGBColor.from_string(INDIGO)
    logo = ASSETS / 'chandigarh-university-logo.jpeg'
    if logo.exists():
        lp = doc.add_paragraph(); lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lp.paragraph_format.space_before = Pt(3); lp.paragraph_format.space_after = Pt(3)
        lp.add_run().add_picture(str(logo), width=Inches(1.55))
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Submitted to:\nMr. Sanjeev Kumar (E20624)')
    r.font.size = Pt(11); r.bold = True; r.font.color.rgb = RGBColor.from_string(DARK)
    doc.add_paragraph(); doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('CHANDIGARH UNIVERSITY\n1st Year • 1st Semester • October 2026')
    r.font.size = Pt(11); r.bold = True; r.font.color.rgb = RGBColor.from_string(NAVY)
    page_break(doc)

    # Index page
    add_heading(doc, 'INDEX', 1, 'StockPilot AI — Python Programming Project Report')
    add_para(doc, 'This index follows the requested project-report format. Page numbers refer to the generated PDF version.')
    index_rows = [
        ('1', 'Project Title and Introduction', '4'),
        ('2', 'Problem Statement and Objectives', '5'),
        ('3', 'Tools and Technologies Used', '6'),
        ('4', 'Algorithm / Methodology', '8'),
        ('5', 'Implementation', '10'),
        ('6', 'Output / Results (Screenshots)', '12'),
        ('7', 'Conclusion', '15'),
        ('A', 'Appendix: Project File Map and References', '16'),
    ]
    add_table(doc, ['No.', 'Content', 'Page'], index_rows, [0.55, 5.8, 0.7], 10)
    add_callout(doc, 'Project at a glance', 'StockPilot AI is a Python Flask inventory application that combines persistent database operations, role-based access, explainable restocking analytics, transaction-safe stock updates, and a responsive dashboard for small businesses.', INDIGO)
    add_para(doc, 'Student: Swarup Chowdhury  |  Course: MCA (AI/ML)  |  Section: 26MAM4(B)  |  Roll No.: 26MCI10229')
    page_break(doc)

    # Student / project overview
    add_heading(doc, 'PROJECT INFORMATION', 1, 'Academic and system profile')
    add_table(doc, ['Student information', 'Details'], [
        ('Name', 'Swarup Chowdhury'), ('Course', 'MCA (AI/ML)'), ('College / University', 'Chandigarh University'),
        ('Section', '26MAM4(B)'), ('Roll number', '26MCI10229'), ('Year / Semester', '1st Year / 1st Semester'),
        ('Subject', 'Python Programming'), ('Submitted to', 'Mr. Sanjeev Kumar (E20624)'),
    ], [2.2, 4.85], 9.5)
    add_heading(doc, 'Project synopsis', 2)
    add_para(doc, 'StockPilot AI is a database-backed inventory and demand-intelligence web application. It is designed for grocery shops, supermarkets, retail stores, wholesalers, and small businesses that need a reliable way to understand stock health and decide what to purchase next.')
    add_para(doc, 'The Python layer is responsible for routing, authentication, session handling, validation, business services, analytics, report generation, and database access. The system is intentionally implemented as a working application instead of a static dashboard: sales reduce inventory, purchase receipts increase inventory, waste uses FEFO batch selection, and reports are generated from stored records.')
    add_callout(doc, 'Important scope note', 'The Preview uses a managed MySQL-compatible database when DATABASE_URL is available and keeps SQLite as a clearly separated local fallback. Oracle SQL and PL/SQL artifacts are included as the academic production target; this report focuses on the Python application layer.', TEAL)
    page_break(doc)

    # 1
    add_heading(doc, '1. Project Title and Introduction', 1)
    add_heading(doc, '1.1 Project title', 2)
    add_para(doc, 'StockPilot AI — Smart Inventory & Demand Intelligence System')
    add_heading(doc, '1.2 Introduction', 2)
    add_para(doc, 'Stock management is a daily operational activity, but many small businesses still rely on notebooks or disconnected spreadsheets. StockPilot AI centralizes product records, suppliers, purchase orders, sales, stock batches, expiry information, alerts, audit history, and explainable recommendations in one application.')
    add_para(doc, 'The application answers three practical questions: What is currently in stock? What needs attention today? What should be purchased next? A dashboard gives an executive view, while Inventory, Sales & Billing, Purchasing, Restocking AI, Expiry & Waste, Analytics, Algorithm Lab, Reports, and Business Assistant pages support detailed work.')
    add_heading(doc, '1.3 Project scope', 2)
    add_bullets(doc, [
        'Secure login, logout, session handling, roles, CSRF checks, and account-scoped workspaces.',
        'Product management with SKU, barcode, category, supplier, price, quantity, reorder level, and expiry fields.',
        'Transactional sales, Decimal-safe totals, stock validation, FEFO deduction, audit movement, and printable/report-ready records.',
        'Purchase orders, receipts, suppliers, low-stock signals, expiry and waste workflows, and CSV exports.',
        'Explainable restocking calculations, baseline forecasting, algorithm demonstrations, and a bounded business assistant.',
    ])
    page_break(doc)

    # 2
    add_heading(doc, '2. Problem Statement and Objectives', 1)
    add_heading(doc, '2.1 Problem statement', 2)
    add_para(doc, 'Manual inventory methods make it difficult to maintain accurate quantities, identify low-stock products, monitor expiry batches, reconcile sales with stock, compare suppliers, and justify purchase decisions. They also increase the risk of duplicate SKUs, incorrect stock deductions, unrecorded waste, and weak auditability.')
    add_para(doc, 'The project solves this problem through a centralized Python web application connected to a normalized database. Every important operation is validated at the backend, written as a transaction, and reflected in the product quantity, stock batch, movement ledger, audit history, and dashboard metrics.')
    add_heading(doc, '2.2 Objectives', 2)
    add_numbered(doc, [
        'Develop a working Python Flask inventory application with a responsive user interface.',
        'Create a persistent data layer for products, suppliers, sales, purchase orders, batches, movements, notifications, users, and audit logs.',
        'Maintain transaction safety so a failed sale, receipt, or waste action does not leave partial data behind.',
        'Implement role-based access for Admin / Owner, Inventory Manager, and Staff users.',
        'Use explainable demand and restocking formulas instead of presenting an unverified black-box prediction.',
        'Demonstrate custom algorithms that have direct business value in inventory search, prioritization, routing, and FEFO handling.',
        'Provide useful reports, onboarding guidance, database persistence, and an application suitable for academic demonstration.',
    ])
    add_callout(doc, 'Success criterion', 'A user should be able to log in, inspect stock health, record a sale, receive a purchase order, review a restocking explanation, record waste, and download a report without manually editing database rows.', INDIGO)
    page_break(doc)

    # 3
    add_heading(doc, '3. Tools and Technologies Used', 1)
    add_table(doc, ['Layer', 'Technology / tool', 'Purpose'], [
        ('Programming language', 'Python 3.11', 'Backend logic, validation, analytics, algorithms, tests, and report/export endpoints.'),
        ('Web framework', 'Flask 3.0.3', 'Routes, sessions, JSON APIs, role checks, request handling, and application shell.'),
        ('Database', 'Managed MySQL-compatible DB', 'Persistent Preview/production records shared across application restarts.'),
        ('Local fallback', 'SQLite', 'Clearly separated local development adapter for offline evaluation.'),
        ('Oracle target', 'Oracle Database + PL/SQL', 'Academic production schema and package artifacts for core operations.'),
        ('Frontend', 'HTML5, CSS3, JavaScript', 'Responsive dashboard, forms, charts, tables, modals, and notifications.'),
        ('Analytics', 'Python Decimal and standard-library calculations', 'Currency-safe totals, demand history, valuation, trends, and restocking signals.'),
        ('Security', 'Hashed passwords, secure sessions, CSRF tokens, role checks', 'Protect accounts and state-changing operations.'),
        ('Testing', 'PyTest and Python compilation', 'Validate algorithms, formulas, error handling, and core workflows.'),
        ('Deployment', 'Docker and managed hosting', 'Package and publish the Flask application with environment-driven configuration.'),
        ('Version control', 'Git', 'Checkpointed project history and reproducible delivery.'),
    ], [1.25, 2.0, 3.7], 8.4)
    add_heading(doc, '3.1 Important Python project files', 2)
    add_table(doc, ['File', 'Responsibility'], [
        ('app.py', 'Flask application, login/registration, protected routes, API responses, report downloads, and health endpoint.'),
        ('db.py', 'Database adapters, schema creation, managed database selection, tenancy migration, and demo seeding.'),
        ('services.py', 'Product creation, sales, purchase receipts, waste, reports, dashboard data, and assistant intents.'),
        ('analytics.py', 'Demand calculations, restocking formulas, category mix, daily sales, movers, and baseline forecast.'),
        ('algorithms.py', 'Custom Merge Sort, Binary Search, hash lookup, Dijkstra, priority queue, and FEFO.'),
        ('templates/index.html', 'Server-rendered application shell and login / new-admin forms.'),
        ('static/app.js and styles.css', 'Client-side view rendering, API actions, charts, forms, responsive styling, and toasts.'),
    ], [2.0, 4.95], 8.7)
    page_break(doc)

    # 4
    add_heading(doc, '4. Algorithm / Methodology', 1)
    add_heading(doc, '4.1 Layered methodology', 2)
    add_numbered(doc, [
        'Authenticate the user and create a protected session containing role and workspace scope.',
        'Read current products, sales, suppliers, batches, purchase orders, and movements from the database.',
        'Validate a business operation in Flask and the service layer before changing any records.',
        'Apply the operation atomically: update product quantity, batch quantity, movement ledger, and audit history together.',
        'Generate dashboard metrics, notifications, recommendations, reports, and assistant answers from stored data.',
    ])
    add_heading(doc, '4.2 Explainable restocking method', 2)
    add_para(doc, 'The restocking engine is a transparent business-rule baseline. It is not presented as a trained machine-learning model. The formulas make each recommendation inspectable by a student, manager, or examiner.')
    add_table(doc, ['Calculation', 'Formula / explanation'], [
        ('Average daily demand', 'Units sold during the lookback period ÷ lookback days.'),
        ('Lead-time demand', 'Average daily demand × supplier lead time in days.'),
        ('Safety stock', 'A conservative buffer based on recent demand; fallback uses reorder-level information.'),
        ('Reorder point', 'Lead-time demand + safety stock.'),
        ('Target stock', 'The larger of a minimum-stock buffer or reorder point plus a planning horizon.'),
        ('Suggested quantity', 'max(0, target stock − current stock).'),
        ('Confidence', 'High, Medium, or Fallback according to available sales history.'),
    ], [1.75, 5.2], 8.8)
    add_heading(doc, '4.3 Algorithms demonstrated', 2)
    add_table(doc, ['Algorithm', 'Business use', 'Complexity'], [
        ('Merge Sort', 'Stable ordering of products, sales, or recommendations.', 'O(n log n) time; O(n) space'),
        ('Binary Search', 'Exact SKU lookup after enforcing sorted input.', 'O(log n); sorted precondition'),
        ('Hash-based lookup', 'Fast SKU or barcode resolution during checkout.', 'O(1) average; O(n) space'),
        ('Dijkstra', 'Shortest route distance for nonnegative warehouse edges.', 'O((V+E) log V)'),
        ('Priority Queue', 'Surface critical restocking alerts first.', 'O(n log n) build-and-pop'),
        ('FEFO', 'Consume the batch with the earliest expiry first.', 'O(n log n) sorting'),
    ], [1.45, 3.75, 1.75], 8.1)
    add_heading(doc, '4.4 FEFO example', 2)
    add_para(doc, 'If a product has batches expiring on 05-Nov, 12-Nov, and 25-Nov, a sale or waste operation first selects the 05-Nov batch. If that batch is insufficient, the remaining quantity is taken from 12-Nov, then 25-Nov. This reduces expiry risk and makes the deduction explainable.')
    page_break(doc)

    # 5
    add_heading(doc, '5. Implementation', 1)
    add_heading(doc, '5.1 Database and persistence', 2)
    add_para(doc, 'The Python application uses an environment-driven adapter. When DATABASE_URL is supplied, the managed MySQL-compatible database is used and data survives process restarts and published deployments. SQLite remains a local fallback. Workspace IDs scope users and operational data so a new admin receives a private empty workspace while the fictional demo dataset remains confined to the demo account.')
    add_para(doc, 'The data model contains users, workspaces, categories, suppliers, products, stock batches, sales, sale items, purchase orders, purchase order items, stock movements, notifications, audit logs, and inventory policies. Foreign-key relationships, unique identifiers, status fields, and validation rules keep the records consistent.')
    add_heading(doc, '5.2 Python route and service flow', 2)
    add_code(doc, '''@app.post("/api/sales")
@roles("admin", "manager", "staff")
def add_sale():
    try:
        return jsonify(process_sale(
            db(), request.get_json(silent=True) or {},
            session["user"]["id"], workspace_id()
        )), 201
    except Exception as exc:
        return jsonify({"error": str(exc)}), 400''')
    add_para(doc, 'The route applies the role and CSRF boundary, then delegates the actual business operation to process_sale. The service validates stock, calculates Decimal-safe totals, selects FEFO batches, updates quantities, writes movement and audit rows, and commits only after all lines succeed. Any exception triggers rollback.')
    add_heading(doc, '5.3 Account onboarding and isolation', 2)
    add_numbered(doc, [
        'The user selects Create new admin on the login page.',
        'The backend validates name, workspace name, unique email, and password length.',
        'A new workspace and owner user are inserted into the managed database.',
        'The user is signed in and sees an empty training dashboard instead of demo records.',
        'All later reads and writes use the workspace ID stored in the protected session.',
    ])
    add_heading(doc, '5.4 Security and validation', 2)
    add_bullets(doc, [
        'Passwords are stored as hashes rather than plain text.',
        'Protected routes require an authenticated session cookie.',
        'State-changing routes require the session CSRF token.',
        'Roles are enforced on the backend, not only by hiding interface controls.',
        'Duplicate SKU and barcode values are rejected by validation and database constraints.',
        'The assistant accepts supported parameterized intents and does not execute arbitrary SQL or code.',
    ])
    add_heading(doc, '5.5 Testing and deployment', 2)
    add_para(doc, 'The project was checked with Python compilation and automated tests covering sorting, binary search, graph validation, alert priority, FEFO selection, restocking calculations, Decimal totals, and insufficient-stock rejection. The application is packaged with Docker and configured to honor the PORT environment variable. The production dependency manifest includes Flask, PyMySQL, pandas, NumPy, oracledb, PyTest, and python-dotenv.')
    page_break(doc)

    # 6 output
    add_heading(doc, '6. Output / Results (Screenshots)', 1)
    add_heading(doc, '6.1 Login and new-admin registration output', 2)
    img = ASSETS / 'stockpilot-login.png'
    if img.exists():
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(img), width=Inches(6.3))
        add_caption(doc, 'Figure 1. StockPilot AI login screen with Sign in and Create new admin options.')
    add_para(doc, 'The login screen presents the demo account for evaluation and a separate new-admin path. The registration flow creates a private workspace rather than copying the demo inventory. This output demonstrates that the Python backend and client-side interface are connected to the account model.')
    page_break(doc)
    add_heading(doc, '6.2 Executive dashboard output', 2)
    img = ASSETS / 'stockpilot-dashboard.png'
    if img.exists():
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(img), width=Inches(6.45))
        add_caption(doc, 'Figure 2. StockPilot AI executive dashboard using database-backed demo records.')
    add_para(doc, 'The dashboard summarizes inventory value, retail value, today’s sales, low-stock count, expiry signals, sales momentum, category mix, critical recommendations, and recent movement. In the captured demo snapshot, the displayed metrics include an inventory value of ₹1,82,370, low stock of 26 products, and 8 expiring signals. These values are presented as fictional demonstration data and are calculated from stored records.')
    add_heading(doc, '6.3 Functional results', 2)
    add_table(doc, ['Tested result', 'Observed behavior'], [
        ('Demo login', 'owner@stockpilot.ai / demo123 opens the Northstar Pantry demo workspace.'),
        ('New-admin registration', 'Creates a private workspace and starts with no operational demo records.'),
        ('Workspace isolation', 'A record created in a new workspace is not visible to the demo account.'),
        ('Product creation', 'Creates a product, opening batch, stock movement, and audit record.'),
        ('Sale processing', 'Validates stock, calculates totals, deducts FEFO batches, and commits as one operation.'),
        ('Purchase receipt', 'Increases stock, creates receipt movement, and updates purchase status.'),
        ('Reports', 'Returns inventory, low-stock, sales, expiry, and audit data as CSV-ready records.'),
        ('Automated tests', 'Four core tests pass after Python compilation.'),
    ], [1.8, 5.15], 8.5)
    page_break(doc)

    # 7 conclusion
    add_heading(doc, '7. Conclusion', 1)
    add_para(doc, 'StockPilot AI demonstrates how Python can be used to build a complete, practical, database-backed business application. The project goes beyond form display: it maintains a stock ledger, protects operations with authentication and roles, preserves data through a managed database, and exposes explainable inventory decisions through analytics and algorithms.')
    add_para(doc, 'The system is useful for an academic Python Programming demonstration because it combines core Python programming concepts with real application concerns: functions and modules, exception handling, Decimal arithmetic, data structures, sorting and searching, graph algorithms, priority queues, database transactions, HTTP routes, validation, report generation, and automated tests.')
    add_para(doc, 'The project can be extended with barcode scanning, supplier price history, richer forecasting models, scheduled notifications, Oracle integration tests, multilingual support, and mobile-friendly receiving workflows. The existing baseline deliberately communicates forecast limitations and keeps business recommendations reviewable by a human user.')
    add_callout(doc, 'Final outcome', 'The delivered application provides a working foundation for small-business inventory intelligence: every recommendation shows its math, every stock-affecting action is auditable, and every new account receives its own persistent workspace.', INDIGO)
    page_break(doc)

    # Appendix
    add_heading(doc, 'APPENDIX: PROJECT FILE MAP AND REFERENCES', 1)
    add_heading(doc, 'A.1 Project file map', 2)
    add_table(doc, ['Area', 'Files / artifacts'], [
        ('Python backend', 'app.py, config.py, db.py, services.py, analytics.py, algorithms.py'),
        ('Frontend', 'templates/index.html, static/app.js, static/styles.css, static/favicon.svg'),
        ('Testing', 'tests/test_core.py'),
        ('Database artifacts', 'oracle/schema.sql, oracle/inventory_pkg.sql, oracle/seed_data.sql'),
        ('Documentation', 'docs/ARCHITECTURE.md, docs/DAA.md, docs/ORACLE_SETUP.md, docs/TRAINING_GUIDE.md, docs/MULTI_USER_PERSISTENCE.md'),
        ('Deployment', 'requirements.txt, Dockerfile, .env.example, public/manus-routes.json'),
    ], [1.7, 5.25], 8.7)
    add_heading(doc, 'A.2 References used for this report', 2)
    add_bullets(doc, [
        'StockPilot AI project source code and database artifacts, Chandigarh University academic project workspace.',
        'StockPilot AI User Training Guide, project documentation.',
        'StockPilot AI Architecture and Algorithm Insights documentation.',
        'Captured StockPilot AI application screenshots from the project Preview and supplied dashboard image.',
    ])
    add_para(doc, 'Prepared for: Python Programming | Chandigarh University | 1st Year, 1st Semester | October 2026')

    path = OUT / 'StockPilot_AI_Python_Programming_Project_Report_Swarup_Chowdhury.docx'
    doc.save(str(path))
    print(path)


if __name__ == '__main__':
    build()
